# -*- coding: utf-8 -*-
"""Builds the Word version of the team coding work distribution from its Markdown source."""
import os, re, sys
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from common import ROOT, REPORTS, FIGS

NAVY, BLUE, INK2 = RGBColor(0x0D, 0x36, 0x6B), RGBColor(0x18, 0x4F, 0x95), RGBColor(0x52, 0x51, 0x4E)
SRC = os.path.join(ROOT, "07_Experimental_Framework", "Team_Coding_Work_Distribution.md")
OUT_DIRS = [REPORTS, os.path.join(ROOT, "07_Experimental_Framework")]


def shade(cell, hex_fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd"); shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), hex_fill)
    tcPr.append(shd)


def para_shade(p, hex_fill):
    pPr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd"); shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), hex_fill)
    pPr.append(shd)


def repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader"); el.set(qn("w:val"), "true"); trPr.append(el)


def add_inline(p, text, size=None, color=None, bold_all=False):
    """Supports **bold**, *italic* and `code`."""
    tokens = re.split(r"(\*\*.+?\*\*|`[^`]+`|(?<![\*\w])\*(?![\s\*]).+?(?<![\s\*])\*(?![\*\w]))", text)
    for t in tokens:
        if not t:
            continue
        if t.startswith("**") and t.endswith("**"):
            r = p.add_run(t[2:-2]); r.bold = True
        elif t.startswith("`") and t.endswith("`"):
            r = p.add_run(t[1:-1]); r.font.name = "Consolas"; r._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
            r.font.color.rgb = RGBColor(0x2C, 0x2C, 0x2A)
        elif t.startswith("*") and t.endswith("*") and len(t) > 2:
            r = p.add_run(t[1:-1]); r.italic = True
        else:
            r = p.add_run(t)
        if bold_all:
            r.bold = True
        if size:
            r.font.size = Pt(size if not (t.startswith("`") and t.endswith("`")) else size - 0.5)
        if color is not None:
            r.font.color.rgb = color


def add_table(doc, rows, width_cm):
    cells = [[c.strip() for c in r.strip().strip("|").split("|")] for r in rows]
    if len(cells) > 1 and all(re.fullmatch(r":?-{2,}:?", c.replace(" ", "")) for c in cells[1] if c):
        cells.pop(1)
    n = max(len(r) for r in cells)
    cells = [r + [""] * (n - len(r)) for r in cells]
    lens = []
    for j in range(n):
        col = [len(cells[i][j]) for i in range(len(cells))]
        lens.append(min(60.0, max(5.0, 0.45 * max(col) + 0.55 * sum(col) / len(col))))
    widths = [max(1.4, width_cm * x / sum(lens)) for x in lens]
    scale = width_cm / sum(widths); widths = [w * scale for w in widths]
    size = 9 if n <= 4 else (8 if n <= 6 else 7.5)
    t = doc.add_table(rows=len(cells), cols=n)
    t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER; t.autofit = False
    for i, r in enumerate(cells):
        for j, v in enumerate(r):
            c = t.cell(i, j); c.width = Cm(widths[j])
            p = c.paragraphs[0]; p.paragraph_format.space_after = Pt(0)
            if i == 0:
                add_inline(p, v, size=size, color=RGBColor(0xFF, 0xFF, 0xFF), bold_all=True); shade(c, "184F95")
            else:
                add_inline(p, v, size=size)
                if i % 2 == 0:
                    shade(c, "F4F4F1")
    repeat_header(t.rows[0])
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def build():
    md = open(SRC, encoding="utf-8").read()
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(1.8); sec.top_margin = sec.bottom_margin = Cm(1.8)
    width_cm = 21.0 - 3.6

    st = doc.styles["Normal"]; st.font.name = "Calibri"; st.font.size = Pt(10.5)
    st._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    for lvl, size, color in [(1, 18, NAVY), (2, 14, BLUE), (3, 12, RGBColor(0x1D, 0x1D, 0x1B))]:
        h = doc.styles[f"Heading {lvl}"]; h.font.name = "Calibri"; h.font.size = Pt(size); h.font.color.rgb = color; h.font.bold = True
        h._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")

    # Title block
    p = doc.add_paragraph(); r = p.add_run("ML ALPHA SIGNAL COMBINATION RESEARCH  |  TEAM PLAN"); r.font.size = Pt(9); r.bold = True; r.font.color.rgb = BLUE
    p = doc.add_paragraph(); r = p.add_run("Team Coding Work Distribution"); r.font.size = Pt(26); r.bold = True; r.font.color.rgb = NAVY
    p = doc.add_paragraph(); r = p.add_run("Absar, Harshit and Maham: three equal workstreams, frozen interface contracts and a conflict-free merge plan")
    r.font.size = Pt(12); r.font.color.rgb = INK2
    p = doc.add_paragraph(); r = p.add_run("Prepared September 2026  |  Companion to PDF 4 (Proposed Research Design) and Workbook 4 (Experiment Matrix)")
    r.font.size = Pt(9); r.italic = True; r.font.color.rgb = INK2

    lines, i, para = md.splitlines(), 0, []
    first_h1 = True

    def flush():
        if para:
            p = doc.add_paragraph(); add_inline(p, " ".join(para)); p.paragraph_format.space_after = Pt(5); para.clear()

    while i < len(lines):
        ln = lines[i]; s = ln.strip()
        if s.startswith("```"):
            flush(); j, buf = i + 1, []
            while j < len(lines) and not lines[j].strip().startswith("```"):
                buf.append(lines[j]); j += 1
            for k, b in enumerate(buf):
                p = doc.add_paragraph(); para_shade(p, "F5F5F3")
                pf = p.paragraph_format; pf.space_after = Pt(0); pf.space_before = Pt(0); pf.line_spacing = 1.0
                r = p.add_run(b if b else " "); r.font.name = "Consolas"; r.font.size = Pt(7.5)
                r._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
            doc.add_paragraph().paragraph_format.space_after = Pt(2)
            i = j + 1; continue
        if not s:
            flush(); i += 1; continue
        if s == "{{PAGEBREAK}}":
            flush(); doc.add_page_break(); i += 1; continue
        if s.startswith("{{FIG:") and s.endswith("}}"):
            flush(); name, _, cap = s[6:-2].partition("|")
            doc.add_picture(os.path.join(FIGS, name.strip()), width=Cm(width_cm))
            doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
            p = doc.add_paragraph(); r = p.add_run(cap.strip()); r.italic = True; r.font.size = Pt(8.5); r.font.color.rgb = INK2
            i += 1; continue
        m = re.match(r"^(#{1,4})\s+(.*)$", s)
        if m:
            flush(); lvl = len(m.group(1))
            if lvl == 1 and first_h1:
                first_h1 = False; i += 1; continue  # title block already added
            doc.add_heading(re.sub(r"\*\*", "", m.group(2)), level=min(lvl, 3)); i += 1; continue
        if s.startswith("|"):
            flush(); rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append(lines[i]); i += 1
            add_table(doc, rows, width_cm); continue
        mb = re.match(r"^(\s*)([-*]|\d+\.)\s+(.*)$", ln)
        if mb:
            flush(); txt = mb.group(3); j = i + 1
            while j < len(lines) and lines[j].strip() and lines[j].startswith("  ") and not re.match(r"^\s*([-*]|\d+\.)\s+", lines[j]):
                txt += " " + lines[j].strip(); j += 1
            style = "List Bullet" if mb.group(2) in "-*" else "List Number"
            p = doc.add_paragraph(style=style); add_inline(p, txt); p.paragraph_format.space_after = Pt(2)
            i = j; continue
        para.append(s); i += 1
    flush()

    # footer page numbers
    fp = sec.footer.paragraphs[0]; fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = fp.add_run()
    for tag, text in [("begin", None), (None, "PAGE"), ("end", None)]:
        if tag:
            el = OxmlElement("w:fldChar"); el.set(qn("w:fldCharType"), tag); run._r.append(el)
        else:
            el = OxmlElement("w:instrText"); el.set(qn("xml:space"), "preserve"); el.text = text; run._r.append(el)
    hp = sec.header.paragraphs[0]; hr = hp.add_run("Team Coding Work Distribution  |  ML Alpha Signal Combination & Portfolio Construction")
    hr.font.size = Pt(8); hr.font.color.rgb = RGBColor(0x89, 0x87, 0x81)

    outs = []
    for d in OUT_DIRS:
        out = os.path.join(d, "Team_Coding_Work_Distribution.docx"); doc.save(out); outs.append(out)
    return outs


if __name__ == "__main__":
    for o in build():
        print(o, round(os.path.getsize(o) / 1e3, 1), "KB")
