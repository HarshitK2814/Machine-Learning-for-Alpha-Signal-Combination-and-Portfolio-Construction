# -*- coding: utf-8 -*-
"""Builds the four Excel workbooks with formulas, filters, conditional formatting, validation and frozen panes."""
import os, re
from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule, DataBarRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from common import (ROOT, REPORTS, J, IDEAS, ACADEMIC, WEIGHTS, EXPERIMENTS, parse_md_table)

HDR_FILL = PatternFill("solid", fgColor="184F95")
HDR_FONT = Font(name="Calibri", bold=True, color="FFFFFF", size=10)
TITLE = Font(name="Calibri", bold=True, size=14, color="0D366B")
SUB = Font(name="Calibri", italic=True, size=9, color="52514E")
BODY = Font(name="Calibri", size=10)
BOLD = Font(name="Calibri", size=10, bold=True)
WRAP = Alignment(wrap_text=True, vertical="top")
CEN = Alignment(horizontal="center", vertical="top", wrap_text=True)
BORDER = Border(bottom=Side(style="thin", color="D9D8D2"))
F_GREEN = PatternFill("solid", fgColor="D6F0DF"); F_AMBER = PatternFill("solid", fgColor="FCEBC7")
F_RED = PatternFill("solid", fgColor="F8D7D7"); F_GREY = PatternFill("solid", fgColor="EEEEEA"); F_BLUE = PatternFill("solid", fgColor="E3EEFB")

ROADMAP = {1: (1, 2), 2: (2, 4), 3: (3, 5), 4: (5, 6), 5: (6, 9), 6: (8, 10), 7: (9, 11), 8: (11, 13), 9: (12, 15), 10: (15, 16)}
PHASE_NAMES = {1: "Literature + research design", 2: "Data infrastructure", 3: "Alpha signal library", 4: "Baseline models",
               5: "ML signal combination", 6: "Portfolio construction", 7: "Robust backtesting", 8: "Robustness + statistical testing",
               9: "Paper writing", 10: "Journal submission"}


def table(ws, headers, rows, widths, r0=1, freeze="B2", center_cols=()):
    for j, h in enumerate(headers, 1):
        c = ws.cell(row=r0, column=j, value=h)
        c.fill, c.font, c.alignment = HDR_FILL, HDR_FONT, Alignment(wrap_text=True, vertical="center")
    for i, r in enumerate(rows, r0 + 1):
        for j, v in enumerate(r, 1):
            c = ws.cell(row=i, column=j, value=v)
            c.font, c.border = BODY, BORDER
            c.alignment = CEN if j in center_cols else WRAP
    for j, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.row_dimensions[r0].height = 34
    last = r0 + len(rows)
    ws.auto_filter.ref = f"A{r0}:{get_column_letter(len(headers))}{last}"
    if freeze:
        ws.freeze_panes = freeze
    return last


def header_block(ws, title, subtitle):
    ws["A1"] = title; ws["A1"].font = TITLE
    ws["A2"] = subtitle; ws["A2"].font = SUB


# ------------------------------------------------------------------ Workbook 1
def wb_literature(refs, path):
    wb = Workbook(); ws = wb.active; ws.title = "Literature"
    headers = ["Paper ID", "Title", "Authors", "Year", "Journal", "Journal Tier", "DOI / URL", "Strand", "Core paper?",
               "Research Question", "Data", "Universe", "Frequency", "Features", "ML Method", "Signal Method", "Portfolio Method",
               "Benchmark", "Evaluation Metrics", "Backtest Period", "Transaction Costs", "Turnover", "Risk Controls",
               "Performance (main result)", "Main Contribution", "Main Limitation", "What Was Not Solved", "Research Gap",
               "Potential Extension", "Relevance to Us (1-5)", "How We Use It / Note", "DOI Verification Status"]
    order = sorted(refs.values(), key=lambda r: (not r.get("core"), -r["relevance"], -r["year"], r["id"]))
    rows, links = [], []
    for r in order:
        g = lambda k: r.get(k, "") or ""
        link = f"https://doi.org/{r['doi']}" if r.get("doi") else g("url")
        links.append(link)
        base = [r["id"], r["title"], r["authors"], r["year"], r["venue"], r["tier"], link, r["category"],
                "Yes" if r.get("core") else "No"]
        if r.get("core"):
            rows.append(base + [g("rq"), g("data"), g("universe"), g("freq"), g("features"), g("ml"), g("signal"),
                                g("portfolio"), g("benchmark"), g("metrics"), g("period"), g("tcost"), g("turnover"), g("risk"),
                                g("result"), g("contribution"), g("limitation"), g("not_solved"), g("gap"), g("extension"),
                                r["relevance"], "Detailed core entry", g("doi_status")])
        else:
            rows.append(base + ["—"] * 14 + [g("result"), "—", "—", "—", "—", "—", r["relevance"], g("limitation"), g("doi_status")])
    widths = [13, 42, 28, 7, 32, 18, 34, 22, 9] + [42] * 20 + [11, 36, 30]
    last = table(ws, headers, rows, widths, freeze="C2", center_cols=(4, 9, 30))
    for i, link in enumerate(links, 2):
        if link:
            c = ws.cell(row=i, column=7); c.hyperlink = link; c.font = Font(name="Calibri", size=10, color="184F95", underline="single")
    ws.conditional_formatting.add(f"AD2:AD{last}", ColorScaleRule(start_type="num", start_value=1, start_color="F4F4F1",
                                                                   end_type="num", end_value=5, end_color="6DA7EC"))
    ws.conditional_formatting.add(f"A2:I{last}", FormulaRule(formula=[f'$I2="Yes"'], fill=F_BLUE))
    ws.conditional_formatting.add(f"AF2:AF{last}", FormulaRule(formula=['ISNUMBER(SEARCH("Crossref",$AF2))'], fill=F_GREEN))
    ws.conditional_formatting.add(f"AF2:AF{last}", FormulaRule(formula=['ISNUMBER(SEARCH("No DOI",$AF2))'], fill=F_AMBER))

    # Summary with live formulas
    sm = wb.create_sheet("Summary")
    header_block(sm, "Literature database summary (live formulas)", "All counts recompute from the Literature sheet; filter there to explore.")
    sm["A4"], sm["B4"] = "Total entries", "=COUNTA(Literature!A:A)-1"
    sm["A5"], sm["B5"] = "Core (detailed) entries", '=COUNTIF(Literature!I:I,"Yes")'
    sm["A6"], sm["B6"] = "Entries with DOI link", '=COUNTIF(Literature!G:G,"https://doi.org/*")'
    sm["A7"], sm["B7"] = "DOIs confirmed/matched via Crossref", '=COUNTIF(Literature!AF:AF,"*Crossref match*")+COUNTIF(Literature!AF:AF,"*Crossref-confirmed*")'
    sm["A8"], sm["B8"] = "Published 2019 or later", '=COUNTIF(Literature!D:D,">=2019")'
    sm["A9"], sm["B9"] = "Relevance 5 entries", '=COUNTIF(Literature!AD:AD,5)'
    for r in range(4, 10):
        sm[f"A{r}"].font = BOLD
    strands = sorted({r["category"] for r in refs.values()})
    sm["A12"], sm["B12"], sm["C12"], sm["D12"] = "Strand", "Entries", "Core entries", "Avg relevance"
    for k, s in enumerate(strands, 13):
        sm[f"A{k}"] = s
        sm[f"B{k}"] = f'=COUNTIF(Literature!$H:$H,$A{k})'
        sm[f"C{k}"] = f'=COUNTIFS(Literature!$H:$H,$A{k},Literature!$I:$I,"Yes")'
        sm[f"D{k}"] = f'=IFERROR(ROUND(AVERAGEIF(Literature!$H:$H,$A{k},Literature!$AD:$AD),2),"")'
    k0 = 13 + len(strands) + 2
    tiers = sorted({r["tier"] for r in refs.values()})
    sm[f"A{k0}"], sm[f"B{k0}"] = "Journal tier", "Entries"
    for k, t in enumerate(tiers, k0 + 1):
        sm[f"A{k}"], sm[f"B{k}"] = t, f'=COUNTIF(Literature!$F:$F,$A{k})'
    k1 = k0 + len(tiers) + 3
    sm[f"A{k1}"], sm[f"B{k1}"], sm[f"C{k1}"], sm[f"D{k1}"] = "Period", "From", "To", "Entries"
    for k, (lab, a, b) in enumerate([("<2000", 1900, 1999), ("2000-2009", 2000, 2009), ("2010-2014", 2010, 2014),
                                     ("2015-2019", 2015, 2019), ("2020-2022", 2020, 2022), ("2023-2026", 2023, 2026)], k1 + 1):
        sm[f"A{k}"], sm[f"B{k}"], sm[f"C{k}"] = lab, a, b
        sm[f"D{k}"] = f'=COUNTIFS(Literature!$D:$D,">="&B{k},Literature!$D:$D,"<="&C{k})'
    for row in (12, k0, k1):
        for col in "ABCD":
            c = sm[f"{col}{row}"]
            if c.value:
                c.fill, c.font = HDR_FILL, HDR_FONT
    sm.column_dimensions["A"].width = 44
    for col in "BCD":
        sm.column_dimensions[col].width = 16

    # Closest prior art
    cp = wb.create_sheet("Closest_Prior_Art")
    rows = parse_md_table(os.path.join(ROOT, "03_Research_Gaps", "Closest_Prior_Art_Tracker.md"), "Tracker")
    hdr, body = rows[0], [[re.sub(r"\*\*|@", "", c) for c in r] for r in rows[1:]]
    last = table(cp, hdr, body, [5, 16, 18, 42, 36, 42, 12, 12], freeze="C2")
    cp.conditional_formatting.add(f"G2:G{last}", FormulaRule(formula=['ISNUMBER(SEARCH("High",$G2))'], fill=F_RED))
    cp.conditional_formatting.add(f"G2:G{last}", FormulaRule(formula=['ISNUMBER(SEARCH("Medium",$G2))'], fill=F_AMBER))

    lg = wb.create_sheet("Legend")
    header_block(lg, "Evidence conventions", "Applies to every text field in the Literature sheet")
    for k, (a, b) in enumerate([("Plain statement", "Verified from abstract, publisher page, authors' repository or search summary of those sources"),
                                ("[inference]", "Research team's reading or synthesis; not a claim by the authors"),
                                ("[verify]", "Could not be confirmed from an accessible primary source; check full text before citing"),
                                ("Relevance 1-5", "Team judgement of relevance to the selected research design"),
                                ("DOI Verification Status", "How the DOI was obtained; see 08_References/crossref_verification.csv")], 4):
        lg[f"A{k}"], lg[f"B{k}"] = a, b; lg[f"A{k}"].font = BOLD; lg[f"B{k}"].alignment = WRAP
    lg.column_dimensions["A"].width = 26; lg.column_dimensions["B"].width = 100
    wb.save(path)


# ------------------------------------------------------------------ Workbook 2
def wb_journals(path):
    wb = Workbook(); ws = wb.active; ws.title = "Journals"
    headers = ["Journal", "Tier", "Publisher", "Scope", "Quant Finance Relevance (1-5)", "ML Relevance (1-5)",
               "Recent Related Papers (IDs)", "Typical Methodology", "Empirical Expectations", "Theory Expectations",
               "Data Expectations", "Code / Data Policy", "Acceptance Difficulty (1-5)", "Fit Score (1-10)",
               "Priority Index (formula)", "Priority Rank (formula)", "Target / Stretch / Backup",
               "OpenAlex 2-yr Mean Citedness (Sep 2026)", "OpenAlex h-index", "Clarivate JIF", "Notes"]
    rows = []
    n = len(J)
    for k, j in enumerate(J, 2):
        rows.append([j["journal"], j["tier"], j["publisher"], j["scope"], j["qf_rel"], j["ml_rel"], j["related"], j["methodology"],
                     j["empirical"], j["theory"], j["data"], j["code_policy"], j["difficulty"], j["fit"],
                     f"=ROUND(0.6*N{k}+0.8*AVERAGE(E{k}:F{k})-0.3*M{k},2)", f"=RANK(O{k},$O$2:$O${n + 1})",
                     j["role"], j["oa_2yr"] if j["oa_2yr"] is not None else "n/a", j["oa_h"] if j["oa_h"] is not None else "n/a",
                     "Not verified - check Clarivate JCR", j["note"]])
    last = table(ws, headers, rows, [30, 22, 22, 44, 11, 11, 30, 36, 36, 30, 26, 38, 11, 10, 11, 10, 22, 13, 10, 18, 40],
                 freeze="B2", center_cols=(5, 6, 13, 14, 15, 16, 18, 19))
    ws.conditional_formatting.add(f"N2:N{last}", ColorScaleRule(start_type="num", start_value=1, start_color="F4F4F1", end_type="num", end_value=10, end_color="6DA7EC"))
    ws.conditional_formatting.add(f"M2:M{last}", ColorScaleRule(start_type="num", start_value=1, start_color="F4F4F1", end_type="num", end_value=5, end_color="EC835A"))
    ws.conditional_formatting.add(f"O2:O{last}", DataBarRule(start_type="min", end_type="max", color="2A78D6"))
    for word, fill in [("Target", F_GREEN), ("Stretch", F_BLUE), ("Backup", F_AMBER), ("Not recommended", F_GREY)]:
        ws.conditional_formatting.add(f"Q2:Q{last}", FormulaRule(formula=[f'ISNUMBER(SEARCH("{word}",$Q2))'], fill=fill, stopIfTrue=True))
    sm = wb.create_sheet("Summary")
    header_block(sm, "Journal strategy summary (live formulas)", "Priority Index = 0.6*Fit + 0.8*mean(QF, ML relevance) - 0.3*Difficulty (team heuristic)")
    sm["A4"], sm["B4"], sm["C4"] = "Role", "Journals", "Avg fit"
    for k, role in enumerate(["Stretch", "Target", "Backup", "Not recommended", "Practitioner", "Early dissemination"], 5):
        sm[f"A{k}"] = role
        sm[f"B{k}"] = f'=COUNTIF(Journals!$Q:$Q,"*{role}*")'
        sm[f"C{k}"] = f'=IFERROR(ROUND(AVERAGEIF(Journals!$Q:$Q,"*{role}*",Journals!$N:$N),2),"")'
    tiers = sorted({j["tier"] for j in J})
    sm["A13"], sm["B13"], sm["C13"], sm["D13"] = "Tier", "Journals", "Avg fit", "Avg difficulty"
    for k, t in enumerate(tiers, 14):
        sm[f"A{k}"] = t
        sm[f"B{k}"] = f'=COUNTIF(Journals!$B:$B,$A{k})'
        sm[f"C{k}"] = f'=IFERROR(ROUND(AVERAGEIF(Journals!$B:$B,$A{k},Journals!$N:$N),2),"")'
        sm[f"D{k}"] = f'=IFERROR(ROUND(AVERAGEIF(Journals!$B:$B,$A{k},Journals!$M:$M),2),"")'
    for row in (4, 13):
        for col in "ABCD":
            if sm[f"{col}{row}"].value:
                sm[f"{col}{row}"].fill, sm[f"{col}{row}"].font = HDR_FILL, HDR_FONT
    sm.column_dimensions["A"].width = 44
    for col in "BCD":
        sm.column_dimensions[col].width = 14
    mn = wb.create_sheet("Metrics_Note")
    header_block(mn, "About the metrics", "Read before quoting any number")
    notes = ["OpenAlex 2-year mean citedness and h-index were retrieved from api.openalex.org/sources in September 2026 (raw JSON in 02_Journal_Strategy).",
             "OpenAlex 2-year mean citedness is an open analogue of an impact factor but is NOT the Clarivate Journal Impact Factor.",
             "Clarivate JIF values were not verified from an accessible primary source and are intentionally left blank.",
             "Chartered ABS AJG and ABDC ratings should be checked on the official lists before use.",
             "Relevance, difficulty and fit scores are team judgements for this specific project, not journal-quality ratings.",
             "Code/data policy entries marked 'Verified' were confirmed on the journal or association policy page."]
    for k, t in enumerate(notes, 4):
        mn[f"A{k}"] = t; mn[f"A{k}"].alignment = WRAP
    mn.column_dimensions["A"].width = 120
    wb.save(path)


# ------------------------------------------------------------------ Workbook 3
def wb_ideas(path):
    wb = Workbook(); wt = wb.active; wt.title = "Weights"
    header_block(wt, "Scoring weights (editable - Ideas sheet recalculates)", "Weights should sum to 1.00")
    wt["A3"], wt["B3"] = "Criterion", "Weight"
    for col in "AB":
        wt[f"{col}3"].fill, wt[f"{col}3"].font = HDR_FILL, HDR_FONT
    for k, (name, w) in enumerate(WEIGHTS, 4):
        wt[f"A{k}"], wt[f"B{k}"] = name, w
    wt["A11"], wt["B11"] = "Sum of weights", "=SUM(B4:B10)"
    wt["A11"].font = BOLD
    wt.conditional_formatting.add("B11", CellIsRule(operator="notEqual", formula=["1"], fill=F_RED))
    dv = DataValidation(type="decimal", operator="between", formula1="0", formula2="1", allow_blank=False)
    wt.add_data_validation(dv); dv.add("B4:B10")
    wt.column_dimensions["A"].width = 42; wt.column_dimensions["B"].width = 12

    ws = wb.create_sheet("Ideas")
    headers = ["Idea ID", "Idea (proposed title)", "Research Question", "Novelty", "Data Feasibility", "Implementation Difficulty",
               "Academic Importance", "Practical Importance", "Publication Potential", "Risk", "Overall Score (formula)",
               "Rank (formula)", "Target Journals", "Hypothesis", "Existing Literature", "What Is Solved", "What Remains Unsolved",
               "Exact Novelty", "Data Required", "Signals Required", "ML Methodology", "Portfolio Methodology", "Benchmark",
               "Experiments", "Evaluation Metrics", "Expected Difficulties"]
    rows = []
    n = len(IDEAS)
    for k, i in enumerate(IDEAS, 2):
        f = (f"=ROUND(D{k}*Weights!$B$4+G{k}*Weights!$B$5+H{k}*Weights!$B$6+I{k}*Weights!$B$7+E{k}*Weights!$B$8"
             f"+(10-F{k})*Weights!$B$9+(10-J{k})*Weights!$B$10,2)")
        rows.append([i["id"], i["title"], i["rq"], i["novelty"], i["data_feas"], i["complexity"], ACADEMIC[i["id"]], i["practical"],
                     i["publication"], i["risk"], f, f"=RANK(K{k},$K$2:$K${n + 1})", i["journals"], i["hypothesis"], i["literature"],
                     i["solved"], i["unsolved"], i["novelty_text"], i["data"], i["signals"], i["ml"], i["portfolio"], i["benchmark"],
                     i["experiments"], i["metrics"], i["difficulties"]])
    last = table(ws, headers, rows, [8, 44, 50] + [11] * 9 + [30] + [40] * 13, freeze="C2", center_cols=tuple(range(4, 13)))
    for col in "DEGHI":
        ws.conditional_formatting.add(f"{col}2:{col}{last}", ColorScaleRule(start_type="num", start_value=0, start_color="F4F4F1", end_type="num", end_value=10, end_color="6DA7EC"))
    for col in "FJ":
        ws.conditional_formatting.add(f"{col}2:{col}{last}", ColorScaleRule(start_type="num", start_value=0, start_color="F4F4F1", end_type="num", end_value=10, end_color="EC835A"))
    ws.conditional_formatting.add(f"K2:K{last}", DataBarRule(start_type="num", start_value=4, end_type="num", end_value=9, color="2A78D6"))
    ws.conditional_formatting.add(f"A2:C{last}", FormulaRule(formula=["$L2=1"], fill=F_GREEN))

    cs = wb.create_sheet("Concept_Selection")
    rows = parse_md_table(os.path.join(ROOT, "04_Paper_Concepts", "Research_Ideas_and_Paper_Concepts.md"), "Selection of the Best Paper Idea")
    hdr = [re.sub(r"\*", "", c) for c in rows[0]]
    body = []
    crit_rows = [r for r in rows[1:] if not r[0].startswith("**Mean")]
    for r in crit_rows:
        body.append([r[0]] + [int(x) for x in r[1:]])
    mean_row = ["Mean (formula)"]
    for j in range(1, len(hdr)):
        col = get_column_letter(j + 1)
        mean_row.append(f"=ROUND(AVERAGE({col}2:{col}{len(body) + 1}),2)")
    body.append(mean_row)
    last = table(cs, hdr, body, [44, 10, 10, 10, 10, 10], freeze="B2", center_cols=(2, 3, 4, 5, 6))
    cs.conditional_formatting.add(f"B2:F{last - 1}", ColorScaleRule(start_type="num", start_value=3, start_color="F4F4F1", end_type="num", end_value=10, end_color="6DA7EC"))
    for col in "ABCDEF":
        cs[f"{col}{last}"].font = BOLD
    wb.move_sheet("Ideas", offset=-1)
    wb.save(path)


# ------------------------------------------------------------------ Workbook 4
def wb_experiments(path):
    wb = Workbook(); ws = wb.active; ws.title = "Experiments"
    headers = ["Experiment ID", "Phase", "Block", "Experiment", "Hypothesis", "Model", "Dataset", "Benchmark", "Metrics",
               "Statistical Test", "Expected Outcome (pre-registered hypothesis, not a result)", "Priority", "Status", "Owner",
               "Planned Start (month)", "Planned End (month)", "Evidence Location", "Notes"]
    rows = []
    for e in EXPERIMENTS:
        eid, phase, block, name, hyp, model, data, bench, metrics, test, exp, prio = e
        a, b = ROADMAP[phase]
        rows.append([eid, phase, block, name, hyp, model, data, bench, metrics, test, exp, prio, "Not started", "", a, b, "", ""])
    last = table(ws, headers, rows, [10, 7, 18, 34, 44, 28, 24, 24, 30, 28, 36, 9, 13, 12, 10, 10, 22, 22], freeze="E2",
                 center_cols=(2, 12, 13, 15, 16))
    dv = DataValidation(type="list", formula1='"Not started,In progress,Blocked,Done,Dropped"', allow_blank=False)
    dv2 = DataValidation(type="list", formula1='"Must,Should,Could"', allow_blank=False)
    ws.add_data_validation(dv); ws.add_data_validation(dv2)
    dv.add(f"M2:M{last}"); dv2.add(f"L2:L{last}")
    for word, fill in [("Done", F_GREEN), ("In progress", F_AMBER), ("Blocked", F_RED), ("Not started", F_GREY)]:
        ws.conditional_formatting.add(f"M2:M{last}", CellIsRule(operator="equal", formula=[f'"{word}"'], fill=fill))
    ws.conditional_formatting.add(f"L2:L{last}", CellIsRule(operator="equal", formula=['"Must"'], font=Font(bold=True, color="0D366B")))

    sm = wb.create_sheet("Summary")
    header_block(sm, "Experiment progress (live formulas)", "Change Status in the Experiments sheet; this sheet updates automatically")
    statuses = ["Not started", "In progress", "Blocked", "Done", "Dropped"]
    blocks = list(dict.fromkeys(e[2] for e in EXPERIMENTS))
    sm["A4"] = "Block"
    for j, s in enumerate(statuses, 2):
        sm.cell(row=4, column=j, value=s)
    sm.cell(row=4, column=7, value="Total"); sm.cell(row=4, column=8, value="% Done")
    for k, b in enumerate(blocks, 5):
        sm.cell(row=k, column=1, value=b)
        for j, s in enumerate(statuses, 2):
            col = get_column_letter(j)
            sm.cell(row=k, column=j, value=f'=COUNTIFS(Experiments!$C:$C,$A{k},Experiments!$M:$M,{col}$4)')
        sm.cell(row=k, column=7, value=f'=COUNTIF(Experiments!$C:$C,$A{k})')
        sm.cell(row=k, column=8, value=f'=IFERROR(E{k}/G{k},0)').number_format = "0%"
    tk = 5 + len(blocks)
    sm.cell(row=tk, column=1, value="All experiments").font = BOLD
    for j in range(2, 8):
        col = get_column_letter(j)
        sm.cell(row=tk, column=j, value=f"=SUM({col}5:{col}{tk - 1})").font = BOLD
    sm.cell(row=tk, column=8, value=f"=IFERROR(E{tk}/G{tk},0)").number_format = "0%"
    for j in range(1, 9):
        c = sm.cell(row=4, column=j); c.fill, c.font = HDR_FILL, HDR_FONT
    sm.conditional_formatting.add(f"H5:H{tk}", DataBarRule(start_type="num", start_value=0, end_type="num", end_value=1, color="1BAF7A"))
    sm.column_dimensions["A"].width = 24
    for j in range(2, 9):
        sm.column_dimensions[get_column_letter(j)].width = 13
    pk = tk + 3
    sm.cell(row=pk, column=1, value="Priority"); sm.cell(row=pk, column=2, value="Count"); sm.cell(row=pk, column=3, value="Done")
    for j in range(1, 4):
        c = sm.cell(row=pk, column=j); c.fill, c.font = HDR_FILL, HDR_FONT
    for k, p in enumerate(["Must", "Should", "Could"], pk + 1):
        sm.cell(row=k, column=1, value=p)
        sm.cell(row=k, column=2, value=f'=COUNTIF(Experiments!$L:$L,$A{k})')
        sm.cell(row=k, column=3, value=f'=COUNTIFS(Experiments!$L:$L,$A{k},Experiments!$M:$M,"Done")')

    fc = wb.create_sheet("Factorial_Cells")
    header_block(fc, "16-cell factorial design with orthogonal contrast coding (formulas)",
                 "Contrast columns (+1/-1) and interactions are formulas; use them to regress monthly net returns of cells on effects (E29)")
    hdr = ["Cell", "Form", "Conditioning", "Objective", "Uncertainty", "Experiment", "x_Nonlinear", "x_State", "x_Economic",
           "x_Uncertainty", "State x Economic", "Nonlinear x Economic", "Nonlinear x State", "Uncertainty x Economic"]
    rows = []
    exp = {("Linear", "Static", "Prediction"): "E20", ("Nonlinear", "Static", "Prediction"): "E21", ("Linear", "State", "Prediction"): "E22",
           ("Nonlinear", "State", "Prediction"): "E23", ("Linear", "Static", "Economic"): "E24", ("Nonlinear", "Static", "Economic"): "E25",
           ("Linear", "State", "Economic"): "E26", ("Nonlinear", "State", "Economic"): "E27"}
    k = 4
    for form in ["Linear", "Nonlinear"]:
        for cond in ["Static", "State"]:
            for obj in ["Prediction", "Economic"]:
                for unc in ["None", "Ensemble shrinkage"]:
                    code = f"{form[0]}-{cond[0] if cond == 'Static' else 'C'}-{obj[0]}-{'0' if unc == 'None' else 'U'}"
                    rows.append([code, form, cond, obj, unc, exp[(form, cond, obj)] + ("" if unc == "None" else " + E28"),
                                 f'=IF(B{k}="Nonlinear",1,-1)', f'=IF(C{k}="State",1,-1)', f'=IF(D{k}="Economic",1,-1)',
                                 f'=IF(E{k}="None",-1,1)', f"=H{k}*I{k}", f"=G{k}*I{k}", f"=G{k}*H{k}", f"=J{k}*I{k}"])
                    k += 1
    table(fc, hdr, rows, [12, 11, 12, 12, 18, 14] + [12] * 8, r0=3, freeze="B4", center_cols=tuple(range(6, 15)))
    fc["A21"] = "Check: every contrast column sums to zero (orthogonal design)"; fc["A21"].font = BOLD
    for j in range(7, 15):
        col = get_column_letter(j)
        fc[f"{col}21"] = f"=SUM({col}4:{col}19)"

    gt = wb.create_sheet("Roadmap_Gantt")
    header_block(gt, "Roadmap Gantt (formula-driven)", "Edit Start/End months; cells with 1 are shaded by conditional formatting")
    gt["A3"], gt["B3"], gt["C3"] = "Phase", "Start month", "End month"
    for m in range(1, 17):
        gt.cell(row=3, column=3 + m, value=m)
    for j in range(1, 20):
        c = gt.cell(row=3, column=j); c.fill, c.font = HDR_FILL, HDR_FONT; c.alignment = CEN
    for k, ph in enumerate(range(1, 11), 4):
        a, b = ROADMAP[ph]
        gt.cell(row=k, column=1, value=f"{ph}. {PHASE_NAMES[ph]}")
        gt.cell(row=k, column=2, value=a); gt.cell(row=k, column=3, value=b)
        for m in range(1, 17):
            col = get_column_letter(3 + m)
            gt.cell(row=k, column=3 + m, value=f'=IF(AND({col}$3>=$B{k},{col}$3<=$C{k}),1,"")').alignment = CEN
    gt.conditional_formatting.add("D4:S13", CellIsRule(operator="equal", formula=["1"], fill=PatternFill("solid", fgColor="2A78D6"),
                                                         font=Font(color="2A78D6")))
    gt.column_dimensions["A"].width = 34
    for j in range(2, 20):
        gt.column_dimensions[get_column_letter(j)].width = 5 if j > 3 else 11
    gt.freeze_panes = "D4"

    ck = wb.create_sheet("Reviewer_Checklist")
    header_block(ck, "Reviewer-proof checklist tracker", "Status dropdown; completion % recalculates")
    lines = open(os.path.join(ROOT, "07_Experimental_Framework", "Reviewer_Proof_Checklist.md"), encoding="utf-8").read().splitlines()
    sec, rows = "", []
    for ln in lines:
        if ln.startswith("## "):
            sec = ln[3:].strip()
        m = re.match(r"^- \[ \] (.*)$", ln.strip())
        if m:
            rows.append([sec, re.sub(r"\[@[^\]]*\]", "", m.group(1)).replace("**", "").strip(), "Not started", "", ""])
    last = table(ck, ["Section", "Checklist item", "Status", "Evidence location", "Owner"], rows, [30, 90, 14, 30, 14], r0=4, freeze="C5",
                 center_cols=(3,))
    dv3 = DataValidation(type="list", formula1='"Not started,In progress,Done,N/A"', allow_blank=False)
    ck.add_data_validation(dv3); dv3.add(f"C5:C{last}")
    ck.conditional_formatting.add(f"C5:C{last}", CellIsRule(operator="equal", formula=['"Done"'], fill=F_GREEN))
    ck.conditional_formatting.add(f"C5:C{last}", CellIsRule(operator="equal", formula=['"In progress"'], fill=F_AMBER))
    ck["D2"] = "Completion:"; ck["D2"].font = BOLD
    ck["E2"] = f'=IFERROR(COUNTIF(C5:C{last},"Done")/(COUNTA(B5:B{last})-COUNTIF(C5:C{last},"N/A")),0)'
    ck["E2"].number_format = "0%"; ck["E2"].font = BOLD
    wb.save(path)


def build_all(refs):
    os.makedirs(REPORTS, exist_ok=True)
    out = {
        "wb1": os.path.join(REPORTS, "Workbook1_Literature_Database.xlsx"),
        "wb2": os.path.join(REPORTS, "Workbook2_Journal_Analysis.xlsx"),
        "wb3": os.path.join(REPORTS, "Workbook3_Research_Idea_Matrix.xlsx"),
        "wb4": os.path.join(REPORTS, "Workbook4_Experiment_Matrix.xlsx"),
    }
    wb_literature(refs, out["wb1"]); wb_journals(out["wb2"]); wb_ideas(out["wb3"]); wb_experiments(out["wb4"])
    return out
